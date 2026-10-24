from __future__ import annotations

import io
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services import (
    draft_state,
    manuscript,
    pdf_export,
    project_service,
    project_state,
    references_io,
    reviewer,
    submission,
)


class SubmissionTestBase(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        root = Path(self.tmp.name)
        self.env = patch.dict(os.environ, {
            "PAPERPILOT_DATA_DIR": str(root / "data"),
            "PAPERPILOT_PROJECT_FILE": str(root / "projects.json"),
            "PAPERPILOT_PROJECT_DIR": str(root / "projects"),
            "PAPERPILOT_AGENT_FILE": str(root / "runs.json"),
            "PAPERPILOT_AGENT_DIR": str(root / "runs"),
        })
        self.env.start()
        self.client = TestClient(app)
        self.pid = project_service.create_project("投稿测试项目", "如何验证投稿包？")["id"]
        memory = project_state.read_memory(self.pid)
        memory["evidence"] = [{
            "id": "pev-1", "paper_id": "P1", "title": "Paper One",
            "section": "Abstract", "page": 1,
            "quote": "存内计算降低数据搬运开销。", "verified": True, "kind": "text",
        }]
        memory["papers"] = {"P1": {
            "id": "P1", "title": "Paper One", "authors": ["Alice Author"],
            "year": 2026, "venue": "Journal of Circuits", "doi": "10.1000/abc", "url": "",
        }}
        project_state.write_memory(self.pid, memory)
        manuscript.update(self.pid, {
            "title": "存内计算 ADC 的低功耗校准方法",
            "subtitle": "课程论文",
            "authors": [{"name": "张三", "affiliation": "微电子学院",
                         "email": "z@example.edu", "orcid": "0000-0002-1825-0097",
                         "corresponding": True}],
            "abstract": "本文研究存内计算 ADC 的校准开销与能效边界。" * 8,
            "keywords": ["存内计算", "ADC", "低功耗", "校准"],
            "funding": "国家自然科学基金",
            "contributions": "张三：概念、方法、写作。",
            "data_availability": "数据可向作者索取。",
            "code_availability": "代码见内部仓库。",
            "competing_interests": "作者声明无利益冲突。",
            "acknowledgments": "感谢实验室同学在仿真上的帮助。",
            "citation_style": "gb7714",
            "language": "zh",
        })
        draft = draft_state.read_draft(self.pid)
        draft["title"] = "存内计算 ADC 的低功耗校准方法"
        draft["citations"] = {"alice2026": {
            "key": "alice2026", "title": "A Low Power ADC", "authors": ["Alice Author"],
            "year": "2026", "venue": "Journal of Circuits", "doi": "10.1000/abc",
            "url": "", "entry_type": "journal", "volume": "12", "issue": "3",
            "pages": "45-67", "publisher": "", "language": "en",
        }}
        draft["sections"] = [{"id": "sec-001", "title": "引言", "blocks": [{
            "id": "blk-001", "type": "paragraph",
            "text": "存内计算可以降低数据搬运开销。",
            "evidence_ids": ["pev-1"], "citation_keys": ["alice2026"],
            "claims": [
                {"text": "存内计算可以降低数据搬运开销", "stance": "supported",
                 "reason": "证据直接支持", "counter_quote": ""},
                {"text": "该方案还可以降低校准成本", "stance": "partial",
                 "reason": "证据只覆盖能耗", "counter_quote": ""},
            ],
        }]}]
        draft_state.write_draft(self.pid, draft)
        self.draft = draft_state.read_draft(self.pid)

    def tearDown(self) -> None:
        self.env.stop()
        self.tmp.cleanup()


class SubmissionServiceTests(SubmissionTestBase):
    def test_profiles_include_blind_and_highlight_limits(self) -> None:
        profiles = {item["id"]: item for item in submission.list_profiles()}
        self.assertEqual(set(profiles), {"course", "generic-journal",
                                         "ieee-conference", "double-blind-journal"})
        self.assertTrue(profiles["double-blind-journal"]["blind"])
        self.assertEqual(profiles["ieee-conference"]["highlights_max"], 0)

    def test_check_adds_profile_rules_and_is_read_only(self) -> None:
        before = Path(self.tmp.name, "projects", self.pid, "draft.json").read_text(encoding="utf-8")
        result = submission.check(self.pid, "generic-journal")
        self.assertIn("profile", result)
        self.assertTrue(result["highlights"])
        self.assertIn("counts", result)
        self.assertEqual(len(result["highlights"]), 1)
        after = Path(self.tmp.name, "projects", self.pid, "draft.json").read_text(encoding="utf-8")
        self.assertEqual(before, after)

    def test_highlights_only_use_supported_claims(self) -> None:
        rows = submission.highlights(self.pid, "course")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["text"], "存内计算可以降低数据搬运开销")
        self.assertNotIn("校准成本", rows[0]["text"])
        self.assertLessEqual(len(rows[0]["text"]), 85)

    def test_long_title_is_a_profile_error(self) -> None:
        manuscript.update(self.pid, {"title": "长" * 160})
        result = submission.check(self.pid, "ieee-conference")
        self.assertTrue(any(item["kind"] == "profile_title" and item["level"] == "error"
                            for item in result["items"]))
        self.assertTrue(result["blocking"])

    def test_cover_letter_contains_title_author_and_highlights(self) -> None:
        letter = submission.cover_letter(self.pid, "course")
        self.assertIn("存内计算 ADC 的低功耗校准方法", letter)
        self.assertIn("张三", letter)
        self.assertIn("z@example.edu", letter)
        self.assertIn("存内计算可以降低数据搬运开销", letter)
        self.assertIn("模板", letter)

    def test_package_contains_the_expected_files(self) -> None:
        with patch.object(pdf_export, "build", return_value=b"%PDF-1.4 test"):
            result = submission.package(self.pid, "double-blind-journal")
        self.assertTrue(result["blinded"])
        self.assertTrue(result["pdf_included"])
        with zipfile.ZipFile(io.BytesIO(result["content"])) as archive:
            names = set(archive.namelist())
            expected = {"manuscript.docx", "manuscript.pdf", "cover-letter.md",
                        "highlights.md", "title-page.md", "checklist.md",
                        "references.ris", "references.bib", "response-template.md", "README.txt"}
            self.assertTrue(expected <= names)
            document = Document(io.BytesIO(archive.read("manuscript.docx")))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertNotIn("张三", text)
            self.assertIn("存内计算 ADC 的低功耗校准方法", text)
            title_page = archive.read("title-page.md").decode("utf-8")
            self.assertIn("张三", title_page)

    def test_package_without_tex_keeps_a_note_instead_of_failing(self) -> None:
        with patch.object(pdf_export, "build", side_effect=pdf_export.PDFUnavailable("no tex")):
            result = submission.package(self.pid, "course")
        self.assertFalse(result["pdf_included"])
        with zipfile.ZipFile(io.BytesIO(result["content"])) as archive:
            self.assertIn("PDF-NOTE.txt", archive.namelist())
            self.assertNotIn("manuscript.pdf", archive.namelist())

    def test_orcid_checksum_and_single_corresponding(self) -> None:
        with self.assertRaises(ValueError):
            manuscript.update(self.pid, {"authors": [
                {"name": "张三", "orcid": "0000-0002-1825-0098", "corresponding": True}
            ]})
        with self.assertRaises(ValueError):
            manuscript.update(self.pid, {"authors": [
                {"name": "张三", "corresponding": True},
                {"name": "李四", "corresponding": True},
            ]})

    def test_title_page_and_cover_letter_include_rich_metadata(self) -> None:
        page = submission.title_page(self.pid)
        self.assertIn("0000-0002-1825-0097", page)
        self.assertIn("通讯作者", page)
        self.assertIn("张三：概念、方法、写作。", page)
        self.assertIn("数据可向作者索取。", page)
        letter = submission.cover_letter(self.pid, "course")
        self.assertIn("作者声明无利益冲突。", letter)
        self.assertIn("ORCID: 0000-0002-1825-0097", letter)

    def test_blinded_package_strips_acknowledgments(self) -> None:
        with patch.object(pdf_export, "build", return_value=b"%PDF-1.4 test"):
            blinded = submission.package(self.pid, "double-blind-journal")
            public = submission.package(self.pid, "course")
        with zipfile.ZipFile(io.BytesIO(blinded["content"])) as archive:
            document = Document(io.BytesIO(archive.read("manuscript.docx")))
            text = " ".join(paragraph.text for paragraph in document.paragraphs)
            self.assertNotIn("感谢实验室", text)
        with zipfile.ZipFile(io.BytesIO(public["content"])) as archive:
            document = Document(io.BytesIO(archive.read("manuscript.docx")))
            text = " ".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("感谢实验室", text)

    def test_unknown_profile_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            submission.check(self.pid, "not-a-profile")


class ReviewerTests(SubmissionTestBase):
    def test_report_is_read_only_and_has_stats(self) -> None:
        before = Path(self.tmp.name, "projects", self.pid, "draft.json").read_text(encoding="utf-8")
        report = reviewer.report(self.pid)
        after = Path(self.tmp.name, "projects", self.pid, "draft.json").read_text(encoding="utf-8")
        self.assertEqual(before, after)
        self.assertIn("stats", report)
        self.assertIn("recommendation", report)
        self.assertEqual(report["stats"]["supported"], 1)
        self.assertEqual(report["stats"]["partial"], 1)

    def test_unsupported_claim_raises_major_revision(self) -> None:
        draft = draft_state.read_draft(self.pid)
        draft["sections"][0]["blocks"][0]["claims"].append({
            "text": "完全无关的断言", "stance": "unsupported",
            "reason": "证据不支持", "counter_quote": "",
        })
        draft_state.write_draft(self.pid, draft)
        report = reviewer.report(self.pid)
        self.assertTrue(report["major"])
        self.assertEqual(report["recommendation"], "major_revision")

    def test_response_template_contains_self_check_and_placeholders(self) -> None:
        text = reviewer.response_template(self.pid)
        self.assertIn("Response to Reviewers", text)
        self.assertIn("自查问题", text)
        self.assertIn("Reviewer comment", text)

    def test_map_comments_ranks_the_relevant_block(self) -> None:
        draft = draft_state.read_draft(self.pid)
        draft["sections"][0]["blocks"].append({
            "id": "blk-002", "type": "paragraph",
            "text": "实验数据集来自三个公开基准，并做了五折交叉验证。",
            "evidence_ids": [], "citation_keys": [], "claims": [],
        })
        draft_state.write_draft(self.pid, draft)
        result = reviewer.map_comments(self.pid, ["请说明实验数据集的来源与划分。"])
        self.assertEqual(result["comments"], 1)
        self.assertEqual(result["items"][0]["matches"][0]["block_id"], "blk-002")
        self.assertTrue(result["items"][0]["suggestion"])

    def test_map_comments_rejects_empty_and_too_many(self) -> None:
        with self.assertRaises(ValueError):
            reviewer.map_comments(self.pid, [])
        with self.assertRaises(ValueError):
            reviewer.map_comments(self.pid, ["x"] * 21)

    def test_response_template_with_comments_includes_comment_and_location(self) -> None:
        text = reviewer.response_template(self.pid, ["请说明实验数据集的来源。"])
        self.assertIn("请说明实验数据集的来源。", text)
        self.assertIn("建议定位", text)


class ReferencesTests(SubmissionTestBase):
    def test_ris_round_trip(self) -> None:
        text = references_io.to_ris(self.draft)
        self.assertIn("TY  - JOUR", text)
        parsed = references_io.parse_ris(text)
        self.assertEqual(parsed[0]["title"], "A Low Power ADC")
        self.assertEqual(parsed[0]["doi"], "10.1000/abc")

    def test_endnote_round_trip(self) -> None:
        text = references_io.to_endnote(self.draft)
        self.assertIn("<records>", text)
        parsed = references_io.parse_endnote(text)
        self.assertEqual(parsed[0]["title"], "A Low Power ADC")
        self.assertEqual(parsed[0]["authors"], ["Alice Author"])

    def test_import_deduplicates_and_bumps_revision(self) -> None:
        payload = references_io.to_ris({"citations": {
            "bob2025": {"key": "bob2025", "title": "Another Paper",
                        "authors": ["Bob Li"], "year": "2025", "venue": "Conf",
                        "doi": "10.1000/xyz", "entry_type": "conference"},
            "dup": {"key": "dup", "title": "A Low Power ADC",
                    "authors": ["Alice Author"], "year": "2026", "venue": "Journal",
                    "doi": "10.1000/abc", "entry_type": "journal"},
        }})
        result = self.client.post(
            f"/api/projects/{self.pid}/draft/references/import",
            json={"format": "ris", "text": payload, "revision": self.draft["revision"]},
        )
        self.assertEqual(result.status_code, 200, result.text)
        body = result.json()
        self.assertEqual(body["added"], 1)
        self.assertEqual(body["skipped"], 1)
        self.assertIn("Li2025", body["keys"])

    def test_import_rejects_stale_revision(self) -> None:
        result = self.client.post(
            f"/api/projects/{self.pid}/draft/references/import",
            json={"format": "ris", "text": "TY  - JOUR\nTI  - X\nER  - \n", "revision": 999},
        )
        self.assertEqual(result.status_code, 409)

    def test_endnote_rejects_dtd(self) -> None:
        with self.assertRaises(references_io.ReferenceValidation):
            references_io.parse_endnote("<!DOCTYPE x [<!ENTITY a 'x'>]><xml/>")

    def test_export_route(self) -> None:
        ris = self.client.get(f"/api/projects/{self.pid}/draft/references/export?format=ris")
        self.assertEqual(ris.status_code, 200)
        self.assertIn("TY  - JOUR", ris.text)
        endnote = self.client.get(f"/api/projects/{self.pid}/draft/references/export?format=endnote")
        self.assertEqual(endnote.status_code, 200)
        self.assertIn("<records>", endnote.text)


class SubmissionApiTests(SubmissionTestBase):
    def test_profiles_and_check(self) -> None:
        response = self.client.get("/api/submission/profiles")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()["profiles"]), 4)
        checked = self.client.post(f"/api/projects/{self.pid}/submission/check",
                                   json={"profile": "course"})
        self.assertEqual(checked.status_code, 200)
        self.assertIn("items", checked.json())

    def test_cover_letter_and_response_template(self) -> None:
        cover = self.client.get(f"/api/projects/{self.pid}/submission/cover-letter?profile=course")
        self.assertEqual(cover.status_code, 200)
        self.assertIn("Cover Letter", cover.text)
        response = self.client.get(f"/api/projects/{self.pid}/reviewer/response-template")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Response to Reviewers", response.text)

    def test_reviewer_report_and_package(self) -> None:
        report = self.client.post(f"/api/projects/{self.pid}/reviewer/report")
        self.assertEqual(report.status_code, 200)
        self.assertIn("recommendation", report.json())
        with patch.object(pdf_export, "build", return_value=b"%PDF-1.4 test"):
            package = self.client.get(f"/api/projects/{self.pid}/submission/package?profile=course")
        self.assertEqual(package.status_code, 200)
        self.assertEqual(package.headers["x-paperpilot-pdf-included"], "1")
        self.assertEqual(package.content[:2], b"PK")

    def test_manuscript_api_round_trips_author_and_submission_fields(self) -> None:
        payload = {
            "authors": [{"name": "张三", "affiliation": "微电子学院",
                         "email": "z@example.edu", "orcid": "0000-0002-1825-0097",
                         "corresponding": True}],
            "contributions": "张三：概念、方法、写作。",
            "data_availability": "数据可向作者索取。",
            "code_availability": "代码见内部仓库。",
            "competing_interests": "作者声明无利益冲突。",
            "acknowledgments": "感谢实验室同学在仿真上的帮助。",
        }
        response = self.client.put(f"/api/projects/{self.pid}/manuscript", json=payload)
        self.assertEqual(response.status_code, 200, response.text)
        stored = self.client.get(f"/api/projects/{self.pid}/manuscript").json()
        self.assertEqual(stored["authors"][0]["orcid"], "0000-0002-1825-0097")
        self.assertTrue(stored["authors"][0]["corresponding"])
        self.assertIn("感谢实验室", stored["acknowledgments"])
        self.assertIn("数据可向作者索取", stored["data_availability"])

    def test_reviewer_mapping_api(self) -> None:
        response = self.client.post(
            f"/api/projects/{self.pid}/reviewer/map-comments",
            json={"comments": ["请说明实验数据集的来源。"]},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["comments"], 1)
        template = self.client.post(
            f"/api/projects/{self.pid}/reviewer/response-template",
            json={"comments": ["请说明实验数据集的来源。"]},
        )
        self.assertEqual(template.status_code, 200)
        self.assertIn("请说明实验数据集的来源。", template.text)
        empty = self.client.post(
            f"/api/projects/{self.pid}/reviewer/map-comments", json={"comments": []}
        )
        self.assertEqual(empty.status_code, 422)

    def test_unknown_project_and_profile(self) -> None:
        self.assertEqual(self.client.get("/api/projects/prj-missing/submission/package").status_code, 404)
        self.assertEqual(self.client.post(
            f"/api/projects/{self.pid}/submission/check", json={"profile": "nope"}
        ).status_code, 422)
