"""DOCX export.

Deliberately a single content model shared with the Markdown/LaTeX exporters:
the draft supplies sections and blocks, ``manuscript`` supplies the front
matter, ``citation_styles`` supplies markers and reference entries. This module
only renders — it never decides what belongs in the paper.

Scope boundaries (see the v5.0 spec): course-paper/bachelor-thesis level
template, no journal templates, no pagination preview, and formulas stay as
LaTeX/Unicode source text because OMML typesetting is explicitly out of scope.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from . import citation_styles

BODY_PT = 10.5
_CAPTION_PREFIXES = ("图", "表", "figure", "fig", "table")


def _set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "SimSun"
    target = run._element.get_or_add_rPr()
    fonts = target.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        target.append(fonts)
    fonts.set(qn("w:eastAsia"), "SimSun")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def _configure_styles(document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(BODY_PT)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:eastAsia"), "SimSun")
    for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        try:
            style = document.styles[name]
        except KeyError:  # pragma: no cover - python-docx ships these
            continue
        style.font.name = "SimSun"
        style.font.size = Pt(size)
        style.font.bold = True


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, end):
        run._r.append(node)
    _set_run_font(run, 9)


def _add_paragraph(document, text: str, *, bold: bool = False, size: float | None = None,
                   center: bool = False, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    _set_run_font(run, size, bold)
    return paragraph


def _is_caption(text: str) -> bool:
    head = str(text or "").strip().lower()
    return head.startswith(_CAPTION_PREFIXES) and any(char.isdigit() for char in head[:8])


def build(draft: dict, manuscript: dict, style: str | None = None, blinded: bool = False) -> bytes:
    style = style or manuscript.get("citation_style") or draft.get("citation_style") or citation_styles.DEFAULT_STYLE
    language = str(manuscript.get("language") or "zh")
    title = str(manuscript.get("title") or draft.get("title") or "研究草稿")
    authors = [] if blinded else [item for item in (manuscript.get("authors") or []) if isinstance(item, dict)]
    abstract = str(manuscript.get("abstract") or "")
    keywords = [str(item) for item in (manuscript.get("keywords") or []) if str(item)]

    document = Document()
    _configure_styles(document)
    if blinded:
        document.core_properties.author = ""
        document.core_properties.last_modified_by = ""

    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    short = title if len(title) <= 30 else title[:30] + "…"
    _set_run_font(header.add_run(short), 9)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number(footer)

    _add_paragraph(document, title, bold=True, size=16, center=True)
    if manuscript.get("subtitle"):
        _add_paragraph(document, str(manuscript["subtitle"]), size=12, center=True)
    if authors:
        names = "，".join(item.get("name", "") for item in authors if item.get("name"))
        _add_paragraph(document, names, center=True)
        affiliations = []
        for item in authors:
            value = str(item.get("affiliation") or "").strip()
            if value and value not in affiliations:
                affiliations.append(value)
        if affiliations:
            _add_paragraph(document, "；".join(affiliations), size=9, center=True)

    if abstract:
        label = "摘要：" if language == "zh" else "Abstract: "
        paragraph = document.add_paragraph()
        _set_run_font(paragraph.add_run(label), BODY_PT, True)
        _set_run_font(paragraph.add_run(abstract), BODY_PT)
    if keywords:
        label = "关键词：" if language == "zh" else "Keywords: "
        separator = "；" if language == "zh" else ", "
        paragraph = document.add_paragraph()
        _set_run_font(paragraph.add_run(label), BODY_PT, True)
        _set_run_font(paragraph.add_run(separator.join(keywords)), BODY_PT)

    for item in draft.get("sections") or []:
        heading = str(item.get("title") or "章节")
        document.add_heading(heading, level=1)
        for block in item.get("blocks") or []:
            text = str(block.get("text") or "").strip()
            if not text:
                continue
            marker = citation_styles.marker(block.get("citation_keys") or [], draft, style)
            block_type = str(block.get("type") or "paragraph")
            if _is_caption(text):
                _add_paragraph(document, text, size=9, center=True)
            elif block_type == "bullet":
                document.add_paragraph(f"- {text}{marker}")
            elif block_type == "quote":
                paragraph = document.add_paragraph(f"{text}{marker}")
                paragraph.paragraph_format.left_indent = Pt(18)
            elif block_type == "note":
                paragraph = document.add_paragraph(f"[注] {text}{marker}")
                paragraph.paragraph_format.left_indent = Pt(18)
            else:
                document.add_paragraph(f"{text}{marker}")

    metadata_sections = [
        ("acknowledgments", "致谢" if language == "zh" else "Acknowledgments", not blinded),
        ("data_availability", "数据可用性" if language == "zh" else "Data availability", True),
        ("code_availability", "代码可用性" if language == "zh" else "Code availability", True),
        ("competing_interests", "利益冲突" if language == "zh" else "Competing interests", True),
    ]
    for field, heading, visible in metadata_sections:
        value = str(manuscript.get(field) or "").strip()
        if value and visible:
            _add_paragraph(document, heading, bold=True, size=10)
            _add_paragraph(document, value, size=9)
    if manuscript.get("funding") and not blinded:
        _add_paragraph(document, ("基金项目：" if language == "zh" else "Funding: ") + str(manuscript["funding"]), size=9)

    rows = citation_styles.references(draft, style)
    if rows:
        document.add_heading("参考文献" if language == "zh" else "References", level=1)
        for row in rows:
            _add_paragraph(document, row["text"], size=9)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
