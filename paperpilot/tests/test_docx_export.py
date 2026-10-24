from __future__ import annotations

import copy
import io
import unittest

from docx import Document

from app.services import docx_export


def _draft() -> dict:
    return {
        "title": "草稿标题",
        "citation_style": "gb7714",
        "sections": [{
            "id": "sec-001", "title": "引言",
            "blocks": [
                {"id": "blk-001", "type": "paragraph",
                 "text": "存内计算降低数据搬运开销", "citation_keys": ["a2024"]},
                {"id": "blk-002", "type": "bullet", "text": "要点一", "citation_keys": []},
                {"id": "blk-003", "type": "paragraph", "text": "图 1 展示了能耗对比。", "citation_keys": []},
                {"id": "blk-004", "type": "paragraph",
                 "text": "目标函数为 $L = x^2$。", "citation_keys": []},
            ],
        }],
        "citations": {"a2024": {
            "key": "a2024", "title": "Paper One", "authors": ["Alice Author"], "year": "2024",
            "venue": "Journal of Things", "doi": "10.1000/xyz", "url": "",
            "entry_type": "journal", "volume": "12", "issue": "3", "pages": "45-67",
            "publisher": "", "language": "en",
        }},
        "revision": 3,
    }


def _meta() -> dict:
    return {
        "title": "低功耗存内计算研究", "subtitle": "课程论文",
        "authors": [{"name": "张三", "affiliation": "微电子学院", "email": "z@example.edu"}],
        "abstract": "本文研究存内计算的能耗问题。",
        "keywords": ["存算一体", "ADC"],
        "funding": "国家自然科学基金",
        "citation_style": "gb7714",
        "language": "zh",
    }


class DocxExportTests(unittest.TestCase):
    def _doc(self, draft=None, meta=None):
        payload = docx_export.build(draft or _draft(), meta or _meta())
        return payload, Document(io.BytesIO(payload))

    def test_payload_is_a_real_docx(self) -> None:
        payload, document = self._doc()
        self.assertTrue(payload.startswith(b"PK"))
        self.assertGreater(len(payload), 5000)
        self.assertTrue(document.paragraphs)

    def test_front_matter_contains_metadata(self) -> None:
        _, document = self._doc()
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("低功耗存内计算研究", text)
        self.assertIn("课程论文", text)
        self.assertIn("张三", text)
        self.assertIn("微电子学院", text)
        self.assertIn("摘要：本文研究存内计算的能耗问题。", text)
        self.assertIn("关键词：存算一体；ADC", text)
        self.assertIn("基金项目：国家自然科学基金", text)

    def test_body_uses_style_markers_and_headings(self) -> None:
        _, document = self._doc()
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("引言", text)
        self.assertIn("存内计算降低数据搬运开销 [1]", text)
        self.assertIn("- 要点一", text)
        headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
        self.assertIn("引言", headings)
        self.assertIn("参考文献", headings)

    def test_caption_and_formula_are_preserved(self) -> None:
        _, document = self._doc()
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("图 1 展示了能耗对比。", text)
        self.assertIn("$L = x^2$", text)

    def test_reference_entry_rendered(self) -> None:
        _, document = self._doc()
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("[1] Alice Author.", text)
        self.assertIn("[J]", text)
        self.assertIn("12(3)", text)
        self.assertIn("45-67", text)

    def test_style_override_changes_markers(self) -> None:
        _, document = self._doc(meta={**_meta(), "citation_style": "apa"})
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("(Author, 2024)", text)
        self.assertIn("Author, A. (2024).", text)

    def test_header_and_footer_are_set(self) -> None:
        _, document = self._doc()
        section = document.sections[0]
        self.assertTrue(section.header.paragraphs[0].text)
        footer_xml = section.footer.paragraphs[0]._p.xml
        self.assertIn("PAGE", footer_xml)

    def test_build_does_not_mutate_inputs(self) -> None:
        draft, meta = _draft(), _meta()
        before = copy.deepcopy((draft, meta))
        docx_export.build(draft, meta)
        self.assertEqual((draft, meta), before)

    def test_english_manuscript_uses_english_labels(self) -> None:
        _, document = self._doc(meta={**_meta(), "language": "en"})
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Abstract:", text)
        self.assertIn("Keywords:", text)
        self.assertIn("References", text)
