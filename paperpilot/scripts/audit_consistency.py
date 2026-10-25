"""Cross-consistency audit: catch "artifacts that do not match their labels".

The product assembles papers, evidence, draft blocks, citations, manuscript metadata and
submission files from several services. A feature can "work" while attaching the wrong
paper, the wrong quote, the wrong title or a stale count to it. This script drives one
offline project end to end and asserts the joins between those parts:

A. agent run     read/compare/topic/ask steps vs their inputs; report legend vs papers;
                 evidence quotes vs the paper text; run papers vs project memory
B. draft         block evidence/citation keys exist; in-text numbering vs reference order;
                 audit issues vs block ids; DOCX text vs block text
C. submission    manuscript title/authors in title page and cover letter (not the project
                 title); package DOCX == draft text; highlights vs supported claims;
                 blinded export removes authors/affiliations/acknowledgments
D. reviewer      stats internal consistency; mapped comments point at real blocks
E. search        source labels match the requested source; layers/years are sane
F. structure     no module calls the model outside llm_gateway (audit/budget/timeout hold)

Usage::

    python scripts/audit_consistency.py [--json out/audit-consistency.json]
"""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TMP = Path(tempfile.mkdtemp(prefix="pp-audit-"))
os.environ.update({
    "PAPERPILOT_DATA_DIR": str(TMP),
    "PAPERPILOT_PROJECT_FILE": str(TMP / "projects.json"),
    "PAPERPILOT_PROJECT_DIR": str(TMP / "projects"),
    "PAPERPILOT_AGENT_FILE": str(TMP / "runs.json"),
    "PAPERPILOT_AGENT_DIR": str(TMP / "runs"),
    "PAPERPILOT_LLM_API_KEY": "",
    "OPENAI_API_KEY": "",
})
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.main import app  # noqa: E402

client = TestClient(app)
RESULTS: list[dict] = []
# Any of these outside llm_gateway.py means the call skips audit, budget,
# timeout and degrade handling — the defect v8.1.3 closed.
GATEWAY_BYPASS_TOKENS = ("llm.chat(", "llm.chat_json(", "import chat
", "import chat_json", "import chat,")


def check(name: str, ok: bool, detail: str = "") -> bool:
    RESULTS.append({"name": name, "ok": bool(ok), "detail": str(detail)[:300]})
    print(("PASS " if ok else "FAIL ") + name + ("  | " + str(detail)[:130] if detail else ""), flush=True)
    return ok


def api(method: str, path: str, payload=None):
    response = client.request(method, path, json=payload)
    if "json" in response.headers.get("content-type", ""):
        return response.status_code, response.json()
    return response.status_code, response.content


def norm(text: str) -> str:
    return " ".join(str(text or "").split())


def docx_text(blob: bytes) -> str:
    document = Document(io.BytesIO(blob))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return norm(" ".join(parts))


def wait_run(run_id: str, timeout: float = 180.0) -> dict:
    import time
    deadline = time.time() + timeout
    while time.time() < deadline:
        _, data = api("GET", f"/api/agent/runs/{run_id}")
        run = (data or {}).get("run") or {}
        if run.get("status") in {"completed", "failed", "cancelled", "awaiting_approval"}:
            return run
        time.sleep(0.4)
    return {}


def main() -> int:
    out_path = None
    parser = argparse.ArgumentParser()
    parser.add_argument("--json", type=Path, default=None)
    args = parser.parse_args()
    out_path = args.json

    # ---------- E. search labels ----------
    status, search = api("POST", "/api/search", {"query": "存算一体 低功耗 ADC", "source": "demo", "limit": 5})
    papers = (search or {}).get("papers") or []
    check("E1 demo 搜索请求与响应来源一致",
          status == 200 and len(papers) == 5 and (search or {}).get("source") == "demo"
          and len({p.get("id") for p in papers}) == len(papers),
          f"code={status} n={len(papers)} source={(search or {}).get('source')}")
    check("E2 每篇题录都带真实来源标记与抽取层级（不是空壳）",
          all(p.get("source") in {"arxiv", "crossref", "europepmc", "doaj", "hal"} for p in papers)
          and all(p.get("extraction_level") for p in papers)
          and all(p.get("title") and p.get("tldr") for p in papers),
          f"levels={sorted({p.get('extraction_level') for p in papers})}")
    check("E3 分层标签合法、年份与标题可用",
          all(p.get("layer") in {"直接相关", "方法参考", "背景脉络"} for p in papers)
          and all(1900 <= int(p.get("year") or 0) <= 2100 for p in papers)
          and len({p.get("title") for p in papers}) == len(papers),
          f"layers={sorted({p.get('layer') for p in papers})}")
    read_body = api("POST", "/api/read", {"paper_id": papers[0]["id"]})[1] or {}
    check("E4 精读返回的正文与请求的论文是同一篇",
          read_body.get("paper", {}).get("id") == papers[0]["id"]
          and len(read_body.get("sections") or []) > 0
          and all(papers[0]["title"][:24] in str(read_body.get("paper", {}).get("title") or "")
                  for _ in [0]),
          f"sections={len(read_body.get('sections') or [])} paper={(read_body.get('paper') or {}).get('id')}")

    # ---------- A. agent run ----------
    status, created = api("POST", "/api/projects", {
        "title": "一致性审计项目", "question": "审计专属研究问题",
        "scope": "CIM ADC", "tags": ["audit"],
    })
    pid = created["project"]["id"]
    status, started = api("POST", "/api/agent/runs", {
        "goal": "梳理存内计算 ADC 的低功耗架构与数字校准开销", "source": "demo",
        "mode": "auto", "project_id": pid, "max_papers": 3, "max_actions": 8,
        "max_llm_calls": 0, "max_seconds": 300,
    })
    run_id = started["run"]["id"]
    run = wait_run(run_id)
    check("A1 Agent auto 运行结束且论文数与预算自洽",
          run.get("status") == "completed" and 2 <= len(run.get("paper_ids") or []) <= 3
          and (run.get("budget_used") or {}).get("actions") <= len(run.get("steps") or []),
          f"status={run.get('status')} papers={len(run.get('paper_ids') or [])} "
          f"steps={len(run.get('steps') or [])} budget={run.get('budget_used')}")
    paper_ids = set(run.get("paper_ids") or [])
    memory_papers = api("GET", f"/api/projects/{pid}/memory?kind=papers&limit=100")[1].get("items") or []
    check("A2 运行的论文集合 == 项目记忆里的论文集合",
          paper_ids == {p.get("id") for p in memory_papers},
          f"run={sorted(paper_ids)} memory={sorted(p.get('id') for p in memory_papers)}")
    titles = {p.get("id"): p.get("title") for p in memory_papers}
    evidence = api("GET", f"/api/projects/{pid}/memory?kind=evidence&limit=100")[1].get("items") or []
    text_by_paper = {}
    for paper in papers:
        body = api("POST", "/api/read", {"paper_id": paper["id"]})[1] or {}
        text_by_paper[paper["id"]] = norm(" ".join(
            s.get("text") or "" for s in body.get("sections") or [])) or norm(body.get("abstract") or "")
    steps = run.get("steps") or []

    expected_prefix = {"search": "检索", "read": "精读", "compare": "横向对比",
                       "topic": "生成候选问题", "ask": "追问", "report": "收束"}
    title_bad = [
        {"tool": step.get("tool"), "title": str(step.get("title") or "")[:30]}
        for step in steps
        if step.get("tool") in expected_prefix
        and not str(step.get("title") or "").startswith(expected_prefix[step["tool"]])
    ]
    check("A2b 每个动作的标题与其工具一致（不会把对比写成精读）",
          not title_bad, json.dumps(title_bad, ensure_ascii=False)[:200])

    misread = []
    for step in steps:
        if step.get("tool") != "read":
            continue
        art = step.get("artifact") or {}
        pid_in = (step.get("input") or {}).get("paper_id")
        art_id = (art.get("paper") or {}).get("id")
        title = str(step.get("title") or "")
        head = str((art.get("paper") or {}).get("title") or "")[:24]
        if art_id != pid_in or not head or head not in title:
            misread.append({"input": pid_in, "artifact": art_id, "title": title[:40], "head": head})
    check("A3 每次精读的输入/产物/步骤标题指向同一篇论文", not misread, json.dumps(misread, ensure_ascii=False)[:200])

    compare_bad = []
    for step in steps:
        if step.get("tool") != "compare":
            continue
        art = step.get("artifact") or {}
        ids_in = set((step.get("input") or {}).get("paper_ids") or [])
        ids_art = set(art.get("paper_ids") or [])
        row_ids = {row.get("id") for row in art.get("rows") or []}
        rows_match = all(str(row.get("title") or "").startswith(str(titles.get(row.get("id")) or "")[:20])
                         for row in art.get("rows") or [])
        if ids_in != ids_art or ids_in != row_ids or not rows_match:
            compare_bad.append({"input": sorted(ids_in), "artifact": sorted(ids_art),
                                "rows": sorted(row_ids), "titles": rows_match})
    check("A4 横向对比的论文集合与每行标题都和所选论文一致",
          not compare_bad, json.dumps(compare_bad, ensure_ascii=False)[:220])

    topic_bad = []
    for step in steps:
        if step.get("tool") != "topic":
            continue
        art = step.get("artifact") or {}
        ids_in = set((step.get("input") or {}).get("paper_ids") or [])
        if set(art.get("paper_ids") or []) != ids_in:
            topic_bad.append("paper_ids")
        for question in art.get("questions") or []:
            for entry in question.get("evidence") or []:
                first = str(entry).split(" · ")[0].strip()
                if first and first not in ids_in and not first.startswith("P"):
                    topic_bad.append(f"{question.get('id')}:{first}")
    check("A5 候选问题只引用本次对比过的论文", not topic_bad, str(topic_bad)[:200])

    value_bad = []
    for step in steps:
        if step.get("tool") != "compare":
            continue
        art = step.get("artifact") or {}
        columns = art.get("columns") or []
        column_keys = {column.get("key") for column in columns}
        # "paper" columns (title/year/venue) are rendered from the row itself, not from values.
        value_keys = {column.get("key") for column in columns if column.get("kind") != "paper"}
        for row in art.get("rows") or []:
            values = row.get("values") or {}
            if set(values) - column_keys:
                value_bad.append(f"{row.get('id')}: extra {sorted(set(values) - column_keys)[:2]}")
            if value_keys - set(values):
                value_bad.append(f"{row.get('id')}: missing {sorted(value_keys - set(values))[:2]}")
    check("A4b 对比表每一行的列都与表头一致（没有错列/缺列）", not value_bad, str(value_bad[:3])[:200])

    question_bad = []
    for step in steps:
        if step.get("tool") != "topic":
            continue
        art = step.get("artifact") or {}
        for question in art.get("questions") or []:
            for entry in question.get("evidence") or []:
                text = str(entry)
                if "：" not in text:
                    continue
                head, quote = text.split("：", 1)
                paper_id = head.split(" · ")[0].strip()
                if re.fullmatch(r"\d{4}\.\d{4,5}v\d", paper_id):
                    if norm(quote)[:24] and norm(quote)[:24] not in text_by_paper.get(paper_id, ""):
                        question_bad.append(f"{question.get('id')}:{paper_id}")
                elif re.match(r"^P\d+\s", text):
                    # Curated questions cite "P3 <title>"; the title must be one of this run's papers.
                    cited = text.split(" ", 1)[1][:20]
                    if not any(str(title or "").startswith(cited) for title in titles.values()):
                        question_bad.append(f"{question.get('id')}:label {cited}")
    check("A5b 候选问题引用的原文能在对应论文里回查", not question_bad, str(question_bad[:3])[:200])

    ask_bad = []
    for step in steps:
        if step.get("tool") != "ask":
            continue
        art = step.get("artifact") or {}
        if str(art.get("question") or "") != str((step.get("input") or {}).get("question") or ""):
            ask_bad.append("question text")
        for cite in art.get("citations") or []:
            cid = cite.get("paper_id")
            if cid not in paper_ids:
                ask_bad.append(f"citation paper {cid}")
            elif str(titles.get(cid) or "")[:24] not in str(cite.get("title") or ""):
                ask_bad.append(f"citation title {cid}")
            quote = norm(cite.get("quote"))
            if quote and quote[:30] not in text_by_paper.get(cid, ""):
                ask_bad.append(f"citation quote {cid}")
    check("A6 追问的题目、引用论文与引用原文都能对上", not ask_bad, str(ask_bad)[:220])

    evidence_bad = []
    for item in evidence:
        if item.get("paper_id") not in paper_ids:
            evidence_bad.append(f"paper {item.get('paper_id')}")
            continue
        quote = norm(item.get("quote"))
        if quote and quote[:30] not in text_by_paper.get(item.get("paper_id"), ""):
            evidence_bad.append(f"quote {item.get('paper_id')}:{quote[:24]}")
    check("A7 每条证据都来自本项目论文且原文可回查", not evidence_bad,
          f"evidence={len(evidence)} issues={evidence_bad[:3]}")

    report_path = Path(str((run.get("report") or {}).get("path") or ""))
    report = report_path.read_text(encoding="utf-8") if report_path.is_file() else ""
    legend = dict(re.findall(r"\[(P\d+)\] = ([^；\n]+)", report))
    used = set(re.findall(r"\[(P\d+)\]", report))
    legend_bad = [key for key in used if key not in legend]
    title_bad = [key for key, name in legend.items()
                 if not any(str(title or "").startswith(name.strip()[:20]) for title in titles.values())]
    check("A8 简报里的 [Pn] 都有图例，且图例指向本次论文",
          bool(legend) and not legend_bad and not title_bad,
          f"legend={len(legend)} used={len(used)} missing={legend_bad[:3]} wrong={title_bad[:3]}")

    refs_section = report.split("## 参考文献")[-1] if "## 参考文献" in report else ""
    ref_markers = set(re.findall(r"^- \[(P\d+)\]", refs_section, re.M))
    check("A8b 参考文献条目带与图例一致的编号",
          bool(legend) and ref_markers == set(legend),
          f"legend={sorted(legend)} refs={sorted(ref_markers)}")

    header = re.search(r"论文：(\d+)/(\d+) · 证据：(\d+)（已核对 (\d+)）", report)
    run_evidence = run.get("evidence") or []
    verified = sum(1 for item in run_evidence if item.get("verified"))
    check("A9 简报开头的动作/论文/证据计数与运行一致",
          bool(header) and int(header.group(1)) == len(paper_ids)
          and int(header.group(2)) == len(paper_ids)
          and int(header.group(3)) == len(run_evidence) and int(header.group(4)) == verified,
          f"header={header.groups() if header else None} papers={len(paper_ids)} "
          f"run_evidence={len(run_evidence)} verified={verified}")

    # ---------- B. draft ----------
    anchored = [item for item in evidence if item.get("verified")][:2]
    draft = api("GET", f"/api/projects/{pid}/draft")[1]["draft"]
    section_title = draft["sections"][0]["title"] if draft.get("sections") else "引言"
    sections = [{
        "id": "sec-001", "title": section_title,
        "blocks": [
            {"id": "blk-001", "type": "paragraph", "status": "checked",
             "text": "审计段落 A：" + norm(anchored[0].get("quote"))[:120], "evidence_ids": [anchored[0]["id"]],
             "citation_keys": [],
             "claims": [{"id": "clm-001", "text": "相似性检测可以减少 SAR 的比较次数。",
                         "stance": "supported", "evidence_ids": [anchored[0]["id"]], "reason": "审计自填"}]},
            {"id": "blk-002", "type": "paragraph", "status": "draft",
             "text": "审计段落 B：该方案在 28nm 下能效提升 987 TOPS/W。", "evidence_ids": [], "citation_keys": [],
             "claims": [{"id": "clm-002", "text": "该方案在 28nm 下能效提升 987 TOPS/W。",
                         "stance": "unsupported", "evidence_ids": [], "reason": "审计自填"}]},
        ],
    }]
    status, saved = api("PUT", f"/api/projects/{pid}/draft", {
        "revision": draft["revision"], "title": "审计草稿", "citation_style": "gb7714", "sections": sections,
    })
    check("B1 草稿保存成功", status == 200, f"code={status}")

    citation = api("POST", f"/api/projects/{pid}/draft/citations",
                   {"paper_id": next(iter(paper_ids)), "revision": api("GET", f"/api/projects/{pid}/draft")[1]["draft"]["revision"]})
    citation_key = (citation[1] or {}).get("key", "")
    current = api("GET", f"/api/projects/{pid}/draft")[1]["draft"]
    current["sections"][0]["blocks"][0]["citation_keys"] = [citation_key]
    cit = next((value for value in (current.get("citations") or {}).values()
                if isinstance(value, dict) and value.get("key") == citation_key), {})
    cit["entry_type"] = "journal"
    cit["volume"] = cit.get("volume") or "12"
    cit["issue"] = cit.get("issue") or "3"
    cit["pages"] = cit.get("pages") or "45-67"
    status, _ = api("PUT", f"/api/projects/{pid}/draft", {
        "revision": current["revision"], "title": current["title"],
        "citation_style": "gb7714", "sections": current["sections"], "citations": current["citations"],
    })
    check("B2 加入一条引用后仍可保存", status == 200 and bool(citation_key), f"code={status} key={citation_key}")

    current = api("GET", f"/api/projects/{pid}/draft")[1]["draft"]
    evidence_ids = {item["id"] for item in evidence}
    citation_keys = {value.get("key") for value in (current.get("citations") or {}).values()
                     if isinstance(value, dict) and value.get("key")}
    bind_bad = []
    for section in current["sections"]:
        for block in section.get("blocks") or []:
            for eid in block.get("evidence_ids") or []:
                if eid not in evidence_ids:
                    bind_bad.append(f"evidence {eid}")
            for key in block.get("citation_keys") or []:
                if key not in citation_keys:
                    bind_bad.append(f"citation {key}")
    check("B3 草稿块的证据与引用键都存在", not bind_bad, str(bind_bad)[:160])

    status, markdown = api("GET", f"/api/projects/{pid}/draft/export?format=markdown")
    md = markdown.decode("utf-8", "replace") if isinstance(markdown, bytes) else str(markdown)
    markers = [int(n) for n in re.findall(r"\[(\d+)\]", md.split("## 参考文献")[0])]
    refs = [line for line in md.split("## 参考文献")[-1].splitlines() if line.strip().startswith("[")]
    check("B4 文内编号按首次出现递增，参考文献条数一致",
          markers == sorted(set(markers)) and markers[:1] == [1] and len(refs) == len(set(markers)),
          f"markers={markers} refs={len(refs)}")
    ref_title = str(cit.get("title") or "")
    check("B5 参考文献条目指向被引论文而不是别的论文",
          bool(refs) and ref_title[:18] in refs[0],
          f"ref={refs[0][:80] if refs else None} title={ref_title[:40]}")

    check("B5b 参考文献没有双句点等机械错位",
          bool(refs) and ".." not in refs[0] and "，，" not in refs[0] and "。。" not in refs[0],
          f"ref={refs[0][:90] if refs else None}")

    status, audit = api("POST", f"/api/projects/{pid}/draft/audit")
    current = api("GET", f"/api/projects/{pid}/draft")[1]["draft"]
    block_ids = {b["id"] for s in current["sections"] for b in s.get("blocks") or []}
    audit_bad = [issue.get("block_id") for issue in (audit or {}).get("issues") or []
                 if issue.get("block_id") and issue["block_id"] not in block_ids]
    check("B6 核查问题指向真实段落", not audit_bad, f"issues={len((audit or {}).get('issues') or [])} bad={audit_bad[:3]}")

    status, docx_blob = api("GET", f"/api/projects/{pid}/draft/export?format=docx")
    text = docx_text(docx_blob) if status == 200 else ""
    missing_blocks = [b["id"] for s in current["sections"] for b in s.get("blocks") or []
                      if norm(b.get("text"))[:24] not in text]
    check("B7 DOCX 正文包含每个草稿段落（内容未错位）",
          status == 200 and not missing_blocks and ref_title[:18] in text,
          f"code={status} missing={missing_blocks} ref_in_docx={ref_title[:18] in text}")

    # ---------- C. manuscript & submission ----------
    author_a, author_b = "审计作者甲", "审计作者乙"
    status, _ = api("PUT", f"/api/projects/{pid}/manuscript", {
        "title": "审计稿件标题", "subtitle": "审计副标题",
        "authors": [{"name": author_a, "affiliation": "审计甲学院", "email": "a@example.org"},
                    {"name": author_b, "affiliation": "审计乙学院", "email": "b@example.org",
                     "orcid": "0000-0002-1825-0097", "corresponding": True}],
        "abstract": "审计摘要：验证标题、作者与投稿件是否一致。",
        "keywords": ["审计", "一致性"], "funding": "审计基金",
        "acknowledgments": "感谢审计助手。", "citation_style": "gb7714", "language": "zh",
    })
    manuscript = api("GET", f"/api/projects/{pid}/manuscript")[1] or {}
    check("C1 稿件字段按写入回读（标题/两位作者/通讯作者）",
          status == 200 and manuscript.get("title") == "审计稿件标题"
          and [a.get("name") for a in manuscript.get("authors") or []] == [author_a, author_b]
          and any(a.get("corresponding") for a in manuscript.get("authors") or []),
          f"code={status} title={manuscript.get('title')} authors={[a.get('name') for a in manuscript.get('authors') or []]}")

    status, early_pkg = api("GET", f"/api/projects/{pid}/submission/package?profile=course")
    page_text = ""
    if status == 200:
        with zipfile.ZipFile(io.BytesIO(early_pkg)) as archive:
            if "title-page.md" in archive.namelist():
                page_text = archive.read("title-page.md").decode("utf-8")
    check("C2 标题页是稿件标题，不是项目标题，且作者/单位/通讯标记齐全",
          status == 200 and "审计稿件标题" in page_text and "一致性审计项目" not in page_text
          and author_a in page_text and author_b in page_text and "审计甲学院" in page_text
          and "通讯" in page_text,
          f"code={status} bytes={len(page_text)}")

    status, cover = api("GET", f"/api/projects/{pid}/submission/cover-letter?profile=course")
    cover_text = cover.decode("utf-8", "replace") if isinstance(cover, bytes) else str(cover)
    check("C3 投稿信用稿件标题与通讯作者，不混入项目标题",
          status == 200 and "审计稿件标题" in cover_text and "一致性审计项目" not in cover_text
          and author_b in cover_text and "课程" in cover_text,
          f"code={status} bytes={len(cover_text)}")

    status, checked = api("POST", f"/api/projects/{pid}/submission/check", {"profile": "course"})
    check("C4 投稿检查的档案与提交的档案一致",
          status == 200 and ((checked or {}).get("profile") or {}).get("id") == "course",
          f"code={status} profile={((checked or {}).get('profile') or {}).get('id')}")
    check("C5 检查项引用的段落都存在",
          all(not str(item.get("target") or "").startswith("blk-") or item["target"] in block_ids
              for item in (checked or {}).get("items") or []),
          f"items={len((checked or {}).get('items') or [])}")

    status, package = api("GET", f"/api/projects/{pid}/submission/package?profile=course")
    names: list[str] = []
    pkg_docx = b""
    pkg_highlights = ""
    pkg_checklist = ""
    if status == 200:
        with zipfile.ZipFile(io.BytesIO(package)) as archive:
            names = archive.namelist()
            pkg_docx = archive.read("manuscript.docx") if "manuscript.docx" in names else b""
            pkg_highlights = archive.read("highlights.md").decode("utf-8") if "highlights.md" in names else ""
            pkg_checklist = archive.read("checklist.md").decode("utf-8") if "checklist.md" in names else ""
    pkg_text = docx_text(pkg_docx) if pkg_docx else ""
    pkg_missing = [b["id"] for s in current["sections"] for b in s.get("blocks") or []
                   if norm(b.get("text"))[:24] not in pkg_text]
    check("C6 投稿包 DOCX 与草稿正文一致（不是空稿或别的项目）",
          status == 200 and "manuscript.docx" in names and not pkg_missing,
          f"code={status} files={len(names)} missing={pkg_missing}")
    supported = [claim for s in current["sections"] for b in s.get("blocks") or []
                 for claim in b.get("claims") or [] if claim.get("stance") == "supported"]
    highlight_lines = [line.lstrip("- ").strip() for line in pkg_highlights.splitlines()
                       if line.strip().startswith("-")]
    check("C7 亮点逐条来自支持档主张（不混入未支持或别的段落）",
          len(highlight_lines) == min(len(supported), 3) and bool(supported)
          and all(any(norm(claim.get("text"))[:12] in norm(line) for claim in supported)
                  for line in highlight_lines),
          f"highlights={highlight_lines[:2]} supported={[c.get('text')[:16] for c in supported]}")
    first_paragraph = next((p.text.strip() for p in Document(io.BytesIO(pkg_docx)).paragraphs
                            if p.text.strip()), "") if pkg_docx else ""
    check("C10 投稿包正文标题用的是稿件标题，不是项目标题或草稿标题",
          first_paragraph == "审计稿件标题"
          and "一致性审计项目" not in pkg_text and "审计草稿" not in pkg_text,
          f"first={first_paragraph[:30]} project_in_body={'一致性审计项目' in pkg_text}")

    highlight_sections = {(row.get("block_id"), row.get("section")) for row in (checked or {}).get("highlights") or []}
    section_by_block = {b["id"]: s.get("title") for s in current["sections"] for b in s.get("blocks") or []}
    check("C11 亮点标注的章节与所在段落一致",
          all(section_by_block.get(block_id) == section for block_id, section in highlight_sections),
          f"rows={sorted(highlight_sections)[:3]}")

    counts_line = re.search(r"错误/提醒/提示：(\d+) / (\d+) / (\d+)", pkg_checklist)
    api_counts = (checked or {}).get("counts") or {}
    check("C12 清单里的错误/提醒/提示计数与检查接口一致",
          bool(counts_line) and tuple(int(value) for value in counts_line.groups())
          == (int(api_counts.get("error") or 0), int(api_counts.get("warn") or 0),
              int(api_counts.get("info") or 0)),
          f"file={counts_line.groups() if counts_line else None} api={api_counts}")

    check("C8 清单声明的是所选档案",
          "课程" in pkg_checklist or "course" in pkg_checklist.lower(), f"bytes={len(pkg_checklist)}")

    status, blinded = api("GET", f"/api/projects/{pid}/submission/package?profile=double-blind-journal")
    blind_docx = b""
    blind_page = ""
    if status == 200:
        with zipfile.ZipFile(io.BytesIO(blinded)) as archive:
            blind_docx = archive.read("manuscript.docx")
            blind_page = archive.read("title-page.md").decode("utf-8")
    blind_text = docx_text(blind_docx) if blind_docx else ""
    check("C9 双盲正文去掉作者、单位、致谢，标题页保留作者",
          status == 200 and author_a not in blind_text and "审计甲学院" not in blind_text
          and "感谢审计助手" not in blind_text and author_a in blind_page,
          f"code={status} author_in_body={author_a in blind_text} ack_in_body={'感谢审计助手' in blind_text}")

    # ---------- D. reviewer ----------
    status, report_json = api("POST", f"/api/projects/{pid}/reviewer/report")
    stats = (report_json or {}).get("stats") or {}
    claim_counts = {"supported": 0, "partial": 0, "unsupported": 0}
    for s in current["sections"]:
        for b in s.get("blocks") or []:
            for claim in b.get("claims") or []:
                stance = claim.get("stance")
                if stance in claim_counts:
                    claim_counts[stance] += 1
    check("D1 审稿视角的支持/部分/不支持计数与草稿主张一致",
          status == 200 and all(int(stats.get(k) or 0) == claim_counts[k] for k in claim_counts),
          f"stats={ {k: stats.get(k) for k in claim_counts} } draft={claim_counts}")
    check("D2 推荐等级与统计口径合法",
          (report_json or {}).get("recommendation") in {"major_revision", "minor_revision", "ready_for_review"}
          and 0 <= float(stats.get("coverage") or 0) <= 1,
          f"recommendation={(report_json or {}).get('recommendation')} coverage={stats.get('coverage')}")

    comments = ["请说明审计数据集的来源与划分。", "请补充与现有方法的定量对比。"]
    status, mapped = api("POST", f"/api/projects/{pid}/reviewer/map-comments", {"comments": comments})
    mapped_bad = [match.get("block_id") for item in (mapped or {}).get("items") or []
                  for match in item.get("matches") or [] if match.get("block_id") not in block_ids]
    check("D3 审稿意见映射到真实段落，条数一致",
          status == 200 and (mapped or {}).get("comments") == len(comments) and not mapped_bad,
          f"code={status} comments={(mapped or {}).get('comments')} bad={mapped_bad[:3]}")

    status, template = api("POST", f"/api/projects/{pid}/reviewer/response-template", {"comments": comments})
    template_text = template.decode("utf-8", "replace") if isinstance(template, bytes) else str(template)
    check("D4 回复模板逐条包含原意见且不串行",
          status == 200 and all(comment in template_text for comment in comments),
          f"code={status} bytes={len(template_text)}")

    fresh = api("GET", f"/api/projects/{pid}/draft")[1]["draft"]
    style_status, style_body = api("PUT", f"/api/projects/{pid}/draft", {
        "revision": fresh["revision"], "title": fresh["title"], "citation_style": "apa",
        "sections": fresh["sections"], "citations": fresh["citations"],
    })
    applied_style = ((style_body or {}).get("draft") or {}).get("citation_style") if isinstance(style_body, dict) else None
    _, apa_blob = api("GET", f"/api/projects/{pid}/draft/export?format=markdown")
    apa_text = apa_blob.decode("utf-8", "replace") if isinstance(apa_blob, bytes) else str(apa_blob)
    apa_refs = [line for line in apa_text.split("## 参考文献")[-1].splitlines() if line.strip()]
    check("B8 切到 APA 后文内标记与参考文献一起变化（不是只换个下拉框）",
          style_status == 200 and "(2026)" in apa_text
          and not (apa_refs[0].startswith("[1]") if apa_refs else True),
          f"code={style_status} stored={applied_style} ref={apa_refs[0][:60] if apa_refs else None}")

    # F. structural invariant: every model call must leave through the gateway, or
    #    audit / budget / timeout / degrade governance silently stops applying.
    bypass: list[str] = []
    for path in sorted((ROOT / "app").rglob("*.py")):
        if path.name in {"llm.py", "llm_gateway.py"}:
            continue
        text = path.read_text(encoding="utf-8")
        hits = [token for token in GATEWAY_BYPASS_TOKENS if token in text]
        if hits:
            bypass.append(f"{path.relative_to(ROOT)}: {hits}")
    check("F1 模型调用只从统一网关出去（否则审计/预算/超时/降级都不生效）",
          not bypass, f"绕过网关：{bypass}")

    passed = sum(1 for item in RESULTS if item["ok"])
    payload = {"checks": RESULTS, "passed": passed, "total": len(RESULTS), "project_id": pid}
    if out_path:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"\n{passed}/{len(RESULTS)} consistency checks passed")
    return 0 if passed == len(RESULTS) else 1


if __name__ == "__main__":
    raise SystemExit(main())
